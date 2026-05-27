import unicodedata
from functools import lru_cache
from openpyxl import load_workbook
from config import EXCEL_PATH


def _norm(s):
    if s is None:
        return ""
    s = str(s).strip().lower()
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode("ascii")
    return s


@lru_cache(maxsize=1)
def load_employees():
    """Read the Excel file and return a list of dicts (header row = keys)."""
    if not EXCEL_PATH.exists():
        return []
    wb = load_workbook(EXCEL_PATH, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    headers = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(next(rows))]
    employees = []
    for row in rows:
        if all(cell is None for cell in row):
            continue
        record = {}
        for h, v in zip(headers, row):
            if hasattr(v, "isoformat"):
                v = v.isoformat()
            record[h] = "" if v is None else v
        employees.append(record)
    return employees


def reload_employees():
    load_employees.cache_clear()
    return load_employees()


def search(query_str, limit=10):
    """Fuzzy-ish match on nom/prenom/matricule."""
    q = _norm(query_str)
    if not q:
        return []
    parts = [p for p in q.split() if p]
    results = []
    for emp in load_employees():
        haystack = " ".join(_norm(emp.get(k, "")) for k in ("matricule", "nom", "prenom", "email"))
        if all(p in haystack for p in parts):
            results.append(emp)
        if len(results) >= limit:
            break
    return results


def get_by_matricule(matricule):
    matricule = str(matricule).strip()
    for emp in load_employees():
        if str(emp.get("matricule", "")).strip() == matricule:
            return emp
    return None
