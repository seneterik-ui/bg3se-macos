"""Generate 5 sample .docx HR templates with {{placeholders}}."""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from config import TEMPLATES_DOCX_DIR


HEADER_COMPANY = "ENTREPRISE EXEMPLE SAS"
HEADER_ADDR = "12 rue de l'Innovation – 75009 Paris – SIRET 123 456 789 00012"


def _new_doc(title):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Header
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(HEADER_COMPANY)
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x2D, 0x31, 0x42)

    p = doc.add_paragraph(HEADER_ADDR)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = Pt(9)
    p.runs[0].font.color.rgb = RGBColor(0x70, 0x70, 0x70)

    doc.add_paragraph()  # spacer

    # Right-aligned date / place
    p = doc.add_paragraph("Fait à {{lieu}}, le {{date_jour}}")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_paragraph()

    # Title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title.upper())
    run.bold = True
    run.font.size = Pt(16)
    run.underline = True

    doc.add_paragraph()
    return doc


def _signature(doc):
    doc.add_paragraph()
    doc.add_paragraph()
    p = doc.add_paragraph("Signature et cachet de l'employeur :")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph("\n\n____________________________").alignment = WD_ALIGN_PARAGRAPH.RIGHT


def attestation_salaire():
    doc = _new_doc("Attestation de salaire")
    doc.add_paragraph(
        "Je soussigné(e), responsable des Ressources humaines de la société "
        f"{HEADER_COMPANY}, atteste que :"
    )
    doc.add_paragraph()
    doc.add_paragraph("{{civilite}} {{nom_complet}}").runs[0].bold = True
    doc.add_paragraph("Matricule : {{matricule}}")
    doc.add_paragraph("Né(e) le : {{date_naissance}}")
    doc.add_paragraph("Demeurant : {{adresse}}")
    doc.add_paragraph()
    doc.add_paragraph(
        "est employé(e) au sein de notre société depuis le {{date_embauche}} "
        "en qualité de {{poste}}, dans le service {{service}}, sous contrat {{type_contrat}}."
    )
    doc.add_paragraph()
    doc.add_paragraph(
        "Son salaire brut mensuel pour la période {{periode}} s'élève à "
    ).add_run("{{salaire_brut}} € brut.").bold = True
    doc.add_paragraph()
    doc.add_paragraph(
        "La présente attestation est délivrée à l'intéressé(e) pour servir et "
        "valoir ce que de droit, notamment auprès de : {{destinataire}}."
    )
    _signature(doc)
    return doc


def demande_arriere():
    doc = _new_doc("Demande de paiement d'arriérés")
    doc.add_paragraph("Émetteur :").runs[0].bold = True
    doc.add_paragraph("{{civilite}} {{nom_complet}} – Matricule {{matricule}}")
    doc.add_paragraph("Service : {{service}} – Poste : {{poste}}")
    doc.add_paragraph()
    doc.add_paragraph("À l'attention du service Paie de {{lieu}}.")
    doc.add_paragraph()
    doc.add_paragraph(
        "Objet : Réclamation de sommes dues au titre du mois de {{mois_concerne}}."
    ).runs[0].bold = True
    doc.add_paragraph()
    doc.add_paragraph("Madame, Monsieur,")
    doc.add_paragraph(
        "Je me permets de revenir vers vous concernant le bulletin de paie du mois de "
        "{{mois_concerne}}. Après vérification, il apparaît qu'une somme de "
    ).add_run("{{montant}} € ").bold = True
    doc.add_paragraph("n'a pas été versée. Motif détaillé :")
    p = doc.add_paragraph("{{motif}}")
    p.paragraph_format.left_indent = Cm(1)
    doc.add_paragraph()
    doc.add_paragraph(
        "Je vous remercie de bien vouloir procéder à la régularisation dans les meilleurs délais "
        "et reste à votre disposition pour tout complément d'information."
    )
    doc.add_paragraph()
    doc.add_paragraph("Cordialement,")
    doc.add_paragraph("{{nom_complet}}")
    return doc


def certificat_travail():
    doc = _new_doc("Certificat de travail")
    doc.add_paragraph(
        f"Nous soussignés, {HEADER_COMPANY}, certifions avoir employé :"
    )
    doc.add_paragraph()
    doc.add_paragraph("{{civilite}} {{nom_complet}}").runs[0].bold = True
    doc.add_paragraph("Matricule : {{matricule}}")
    doc.add_paragraph("Demeurant : {{adresse}}")
    doc.add_paragraph("N° de Sécurité sociale : {{numero_ss}}")
    doc.add_paragraph()
    doc.add_paragraph(
        "du {{date_embauche}} au {{date_fin}}, en qualité de {{poste}} "
        "(contrat {{type_contrat}}), au sein du service {{service}}."
    )
    doc.add_paragraph()
    doc.add_paragraph("Motif de fin de contrat : {{motif_fin}}.")
    doc.add_paragraph()
    doc.add_paragraph(
        "Conformément à l'article L1234-19 du Code du travail, "
        "l'intéressé(e) est libre de tout engagement vis-à-vis de notre société."
    )
    doc.add_paragraph(
        "En foi de quoi, le présent certificat lui est délivré pour servir et "
        "valoir ce que de droit."
    )
    _signature(doc)
    return doc


def demande_mutuelle():
    doc = _new_doc("Prise en charge santé / mutuelle")
    doc.add_paragraph("Demandeur :").runs[0].bold = True
    doc.add_paragraph("{{civilite}} {{nom_complet}}")
    doc.add_paragraph("Matricule : {{matricule}} – N° SS : {{numero_ss}}")
    doc.add_paragraph("Service : {{service}} – Poste : {{poste}}")
    doc.add_paragraph()
    doc.add_paragraph("Type de demande : ").add_run("{{type_demande}}").bold = True
    doc.add_paragraph()
    doc.add_paragraph("Ayants droit concernés :")
    p = doc.add_paragraph("{{ayants_droit}}")
    p.paragraph_format.left_indent = Cm(1)
    doc.add_paragraph()
    doc.add_paragraph("Précisions complémentaires :")
    p = doc.add_paragraph("{{precisions}}")
    p.paragraph_format.left_indent = Cm(1)
    doc.add_paragraph()
    doc.add_paragraph(
        "Je vous remercie de bien vouloir prendre en compte ma demande "
        "et reste à disposition pour fournir toute pièce justificative."
    )
    doc.add_paragraph()
    doc.add_paragraph("Signature du salarié :").alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph("\n____________________________").alignment = WD_ALIGN_PARAGRAPH.RIGHT
    return doc


def attestation_juridique():
    doc = _new_doc("Attestation employeur")
    doc.add_paragraph(
        f"Je soussigné(e), représentant de la société {HEADER_COMPANY}, "
        "atteste sur l'honneur les éléments suivants concernant :"
    )
    doc.add_paragraph()
    doc.add_paragraph("{{civilite}} {{nom_complet}}").runs[0].bold = True
    doc.add_paragraph("Matricule interne : {{matricule}}")
    doc.add_paragraph("Né(e) le : {{date_naissance}}")
    doc.add_paragraph()
    doc.add_paragraph(
        "Salarié(e) en contrat {{type_contrat}} depuis le {{date_embauche}}, "
        "occupant le poste de {{poste}} au sein du service {{service}}."
    )
    doc.add_paragraph()
    doc.add_paragraph("Objet de l'attestation : ").add_run("{{objet}}").bold = True
    doc.add_paragraph("Destinataire : {{destinataire}}")
    doc.add_paragraph()
    doc.add_paragraph("Éléments à attester :")
    p = doc.add_paragraph("{{details}}")
    p.paragraph_format.left_indent = Cm(1)
    doc.add_paragraph()
    doc.add_paragraph(
        "La présente attestation est délivrée à l'intéressé(e) pour faire valoir "
        "ses droits, conformément aux dispositions du Code du travail."
    )
    _signature(doc)
    return doc


GENERATORS = {
    "attestation_salaire.docx": attestation_salaire,
    "demande_arriere.docx": demande_arriere,
    "certificat_travail.docx": certificat_travail,
    "demande_mutuelle.docx": demande_mutuelle,
    "attestation_juridique.docx": attestation_juridique,
}


def main():
    TEMPLATES_DOCX_DIR.mkdir(parents=True, exist_ok=True)
    for filename, gen in GENERATORS.items():
        path = TEMPLATES_DOCX_DIR / filename
        doc = gen()
        doc.save(path)
        print(f"Gabarit créé : {path}")


if __name__ == "__main__":
    main()
