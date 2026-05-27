"""Generate a sample employees.xlsx with 12 fictional employees."""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from config import EXCEL_PATH

HEADERS = [
    "matricule", "nom", "prenom", "date_naissance", "sexe",
    "email", "telephone", "adresse",
    "date_embauche", "type_contrat", "poste", "service",
    "salaire_brut", "manager", "numero_ss", "iban",
]

EMPLOYEES = [
    ["E001", "Martin",   "Sophie",   "1985-03-12", "F", "sophie.martin@entreprise.fr",   "0612345678", "12 rue de Rivoli, 75001 Paris",       "2018-09-01", "CDI", "Comptable",            "Finance",       3200.00, "Jean Dupont",       "2850375120458", "FR7630006000011234567890189"],
    ["E002", "Dubois",   "Lucas",    "1990-07-22", "M", "lucas.dubois@entreprise.fr",    "0623456789", "5 avenue Foch, 75116 Paris",          "2020-01-15", "CDI", "Développeur",          "IT",            3800.00, "Claire Bernard",    "1900775120459", "FR7630006000011234567890190"],
    ["E003", "Bernard",  "Claire",   "1982-11-05", "F", "claire.bernard@entreprise.fr", "0634567890", "27 boulevard Voltaire, 75011 Paris",  "2015-06-01", "CDI", "Lead Tech",            "IT",            5200.00, "Pierre Garnier",    "2821175120460", "FR7630006000011234567890191"],
    ["E004", "Petit",    "Thomas",   "1995-02-18", "M", "thomas.petit@entreprise.fr",   "0645678901", "8 rue Lafayette, 75009 Paris",        "2023-03-20", "CDD", "Assistant marketing",  "Marketing",     2400.00, "Sophie Martin",     "1950275120461", "FR7630006000011234567890192"],
    ["E005", "Garnier",  "Pierre",   "1975-09-30", "M", "pierre.garnier@entreprise.fr","0656789012", "44 rue Saint-Honoré, 75001 Paris",    "2010-04-12", "CDI", "Directeur Général",    "Direction",     7800.00, "",                  "1750975120462", "FR7630006000011234567890193"],
    ["E006", "Lefevre",  "Camille",  "1992-12-08", "F", "camille.lefevre@entreprise.fr","0667890123", "15 rue du Bac, 75007 Paris",          "2021-11-02", "CDI", "Chef de projet",       "IT",            4100.00, "Claire Bernard",    "2921275120463", "FR7630006000011234567890194"],
    ["E007", "Roux",     "Antoine",  "1988-05-14", "M", "antoine.roux@entreprise.fr",   "0678901234", "9 avenue Montaigne, 75008 Paris",     "2019-08-26", "CDI", "Commercial Senior",    "Ventes",        3500.00, "Pierre Garnier",    "1880575120464", "FR7630006000011234567890195"],
    ["E008", "Moreau",   "Julie",    "1993-04-27", "F", "julie.moreau@entreprise.fr",   "0689012345", "33 rue Saint-Dominique, 75007 Paris", "2022-05-10", "CDI", "Designer UX",          "Produit",       3300.00, "Camille Lefevre",   "2930475120465", "FR7630006000011234567890196"],
    ["E009", "Faure",    "Nicolas",  "1980-08-19", "M", "nicolas.faure@entreprise.fr",  "0690123456", "21 rue de Vaugirard, 75006 Paris",    "2016-02-08", "CDI", "Responsable RH",       "Ressources humaines", 4600.00, "Pierre Garnier","1800875120466", "FR7630006000011234567890197"],
    ["E010", "Mercier",  "Élodie",   "1996-10-11", "F", "elodie.mercier@entreprise.fr", "0601234567", "7 rue Cler, 75007 Paris",             "2024-06-03", "CDD", "Chargée de communication","Marketing",   2600.00, "Sophie Martin",     "2961075120467", "FR7630006000011234567890198"],
    ["E011", "Blanc",    "Marc",     "1970-01-25", "M", "marc.blanc@entreprise.fr",     "0612000000", "55 avenue de la République, 75011 Paris","2008-09-15","CDI","Architecte logiciel", "IT",            5800.00, "Claire Bernard",    "1700175120468", "FR7630006000011234567890199"],
    ["E012", "Henry",    "Aurélie",  "1989-06-03", "F", "aurelie.henry@entreprise.fr",  "0623000000", "18 rue de Sèvres, 75006 Paris",       "2017-10-23", "CDI", "Juriste",              "Juridique",     4200.00, "Nicolas Faure",     "2890675120469", "FR7630006000011234567890200"],
]


def main():
    EXCEL_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Employes"
    ws.append(HEADERS)
    header_fill = PatternFill("solid", fgColor="2D3142")
    header_font = Font(bold=True, color="FFFFFF")
    for col_idx, _ in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for emp in EMPLOYEES:
        ws.append(emp)
    widths = [10, 14, 14, 14, 6, 32, 14, 38, 14, 8, 26, 22, 12, 22, 16, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"
    wb.save(EXCEL_PATH)
    print(f"Fichier Excel créé : {EXCEL_PATH} ({len(EMPLOYEES)} employés)")


if __name__ == "__main__":
    main()
